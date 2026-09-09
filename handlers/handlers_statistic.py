import asyncio
import calendar
import os
import tempfile
from datetime import datetime, date, timedelta
import openpyxl
from aiogram import Router
from aiogram.filters import Command
from aiogram.types import Message, FSInputFile
from openpyxl.styles import Alignment, Border, Font, Side, PatternFill
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from sqlalchemy import select, func

from bot import sql
from config import ADMIN_IDS, CHECKER_IDS
from logging_config import logger
from config_bd.models import AsyncSessionLocal, Users, Payments, PaymentsStars, PaymentsCryptobot, PaymentsCards, \
    PaymentsPlategaCrypto, PaymentsWataSBP, PaymentsWataCard, PaymentsFkSBP, PlategaRecurent

router = Router()

# Тестовые / служебные суммы — не выводим в разбивке по номиналам.
_EXCLUDED_PAYMENT_AMOUNTS = {1, 10, 50}

# Успешные статусы во всех таблицах оплат (confirmed / paid / CONFIRMED).
_PAYMENT_OK_STATUSES = ("confirmed", "paid")

_ANAL_PAY_MONTHS = (6, 7, 8)
_ANAL_PAY_MONTH_RU = {6: "Июнь", 7: "Июль", 8: "Август"}


# ---------- Вспомогательные функции конвертации ----------
def convert_stars_to_rub(amount: int) -> int:
    """1 звезда Telegram = 1 условный рубль в отчётах."""
    return amount


def _stars_amount_to_rub(amount) -> int | None:
    if amount is None:
        return None
    return convert_stars_to_rub(int(amount))


def _normalize_payment_amount(amount) -> int | None:
    """Приводит сумму платежа к целым рублям; None — если сумма служебная/нулевая."""
    try:
        rub = int(round(float(amount)))
    except (TypeError, ValueError):
        return None
    if rub <= 0 or rub in _EXCLUDED_PAYMENT_AMOUNTS:
        return None
    return rub


async def _discover_payment_amounts(session) -> list[int]:
    """Все уникальные суммы подтверждённых платежей по таблицам, кроме 1/10/50."""
    amounts: set[int] = set()

    sources = [
        (Payments, Payments.status == 'confirmed', Payments.amount),
        (PaymentsCards, PaymentsCards.status == 'confirmed', PaymentsCards.amount),
        (PaymentsPlategaCrypto, PaymentsPlategaCrypto.status == 'confirmed', PaymentsPlategaCrypto.amount),
        (PaymentsWataSBP, PaymentsWataSBP.status == 'confirmed', PaymentsWataSBP.amount),
        (PaymentsWataCard, PaymentsWataCard.status == 'confirmed', PaymentsWataCard.amount),
        (PaymentsFkSBP, PaymentsFkSBP.status == 'confirmed', PaymentsFkSBP.amount),
        (PaymentsStars, PaymentsStars.status == 'confirmed', PaymentsStars.amount),
        (PaymentsCryptobot, PaymentsCryptobot.status == 'paid', PaymentsCryptobot.amount),
    ]
    for _model, status_filter, amount_col in sources:
        stmt = select(amount_col).where(status_filter).distinct()
        for (raw,) in (await session.execute(stmt)).all():
            rub = _normalize_payment_amount(raw)
            if rub is not None:
                # Stars уже в «условных рублях» (1:1), cryptobot — тоже в рублях.
                amounts.add(rub)
    return sorted(amounts)


class PaymentRecord:
    """Унифицированная запись о платеже."""
    def __init__(self, amount: int, is_gift: bool, time_created: datetime):
        self.amount = amount
        self.is_gift = is_gift
        self.time_created = time_created


def _sync_build_analytics_excel(
    monthly_data: dict,
    daily_data_by_month: dict,
    payment_amounts: list[int],
) -> str:
    # --- Создание Excel файла ---
    wb = openpyxl.Workbook()
    ws_main = wb.active
    ws_main.title = "Помесячная аналитика"

    headers = ['Показатель'] + list(monthly_data.keys())
    ws_main.append(headers)

    metric_rows = [
        ('Новые пользователи (всего)', 'new_total'),
        ('Новые пользователи (залив)', 'new_zaliv'),
        ('Новые пользователи (сарафан)', 'new_saraf'),
        ('Взяли ключ (всего)', 'key_total'),
        ('Взяли ключ (залив)', 'key_zaliv'),
        ('Взяли ключ (сарафан)', 'key_saraf'),
        ('Подключились (всего)', 'connect_total'),
        ('Подключились (залив)', 'connect_zaliv'),
        ('Подключились (сарафан)', 'connect_saraf'),
        ('Платежи новых (сумма, всего)', 'pay_new_sum_total'),
        ('Платежи новых (уникальных, всего)', 'pay_new_users_total'),
        ('Платежи новых (сумма, залив)', 'pay_new_sum_zaliv'),
        ('Платежи новых (уникальных, залив)', 'pay_new_users_zaliv'),
        ('Платежи новых (сумма, сарафан)', 'pay_new_sum_saraf'),
        ('Платежи новых (уникальных, сарафан)', 'pay_new_users_saraf'),
        ('Общая выручка (₽)', 'total_revenue'),
        ('Количество платежей', 'total_payments'),
        ('AOV (₽)', 'aov'),
        ('ARPU (₽)', 'arpu'),
        ('Пользователей на конец месяца', 'cumulative_users'),
    ]
    for amt in payment_amounts:
        metric_rows.append((f'Платежей {amt}₽ (шт)', f'sum_{amt}_count'))
        metric_rows.append((f'Сумма {amt}₽ (₽)', f'sum_{amt}_amount'))
    metric_rows.extend([
        ('Подарков (шт)', 'gift_count'),
        ('Сумма подарков (₽)', 'gift_amount'),
    ])

    row_idx = 2
    for label, key in metric_rows:
        row = [label]
        ws_main.append(row)
        col_idx = 2
        for month in monthly_data.keys():
            value = monthly_data[month].get(key, 0)
            if key in ('aov', 'arpu'):
                cell_value = round(value, 2)
            else:
                cell_value = value if isinstance(value, int) else round(value, 2)
            ws_main.cell(row=row_idx, column=col_idx, value=cell_value)
            col_idx += 1
        row_idx += 1

    # Оформление
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    light_green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    light_red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    for cell in ws_main[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    month_columns = list(monthly_data.keys())
    for r in range(2, row_idx):
        for c in range(1, ws_main.max_column + 1):
            ws_main.cell(row=r, column=c).border = thin_border
        jan_cell = ws_main.cell(row=r, column=2)
        jan_cell.fill = yellow_fill
        for col_idx in range(3, 2 + len(month_columns)):
            current = ws_main.cell(row=r, column=col_idx)
            prev = ws_main.cell(row=r, column=col_idx-1)
            try:
                cur_val = float(current.value)
                prev_val = float(prev.value)
            except (TypeError, ValueError):
                continue
            if cur_val > prev_val:
                current.fill = light_green_fill
            elif cur_val < prev_val:
                current.fill = light_red_fill

    for col in ws_main.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws_main.column_dimensions[col_letter].width = min(max_len + 2, 50)

    ws_main.freeze_panes = 'B2'

    # Листы по месяцам с графиками
    for month_key, daily_data in daily_data_by_month.items():
        ws = wb.create_sheet(title=month_key[:31])
        ws.append(['День', 'Новые', 'Взяли ключ', 'Подключились', 'Платили',
                   'Всего пользователей (накопительно)', 'Всего ключей (накопительно)', 'Всего подключений (накопительно)'])
        for d in daily_data:
            ws.append([
                d['day'],
                d['new'],
                d['key'],
                d['connect'],
                d['paid'],
                d['cum_users'],
                d['cum_key'],
                d['cum_connect']
            ])

        for row in ws.iter_rows(min_row=1, max_row=len(daily_data)+1, min_col=1, max_col=8):
            for cell in row:
                cell.border = thin_border

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 20)

        # Линейный график (накопительные)
        chart1 = LineChart()
        chart1.title = "Накопительные показатели"
        chart1.style = 13
        chart1.y_axis.title = "Количество"
        chart1.x_axis.title = "День месяца"
        data = Reference(ws, min_col=6, max_col=8, min_row=1, max_row=len(daily_data)+1)
        dates = Reference(ws, min_col=1, min_row=2, max_row=len(daily_data)+1)
        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(dates)
        if len(chart1.series) >= 3:
            chart1.series[0].graphicalProperties.line.solidFill = "0000FF"
            chart1.series[1].graphicalProperties.line.solidFill = "00B0F0"
            chart1.series[2].graphicalProperties.line.solidFill = "000000"
        ws.add_chart(chart1, "J2")

        # Столбцовая диаграмма (ежедневные)
        chart2 = BarChart()
        chart2.title = "Ежедневные показатели"
        chart2.style = 13
        chart2.y_axis.title = "Количество"
        chart2.x_axis.title = "День месяца"
        data2 = Reference(ws, min_col=2, max_col=5, min_row=1, max_row=len(daily_data)+1)
        chart2.add_data(data2, titles_from_data=True)
        chart2.set_categories(dates)
        ws.add_chart(chart2, "J20")

    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _sync_build_anal_payment_excel(year: int, daily_by_month: dict) -> str:
    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Данные"

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    ws_data.append(
        ["День"] + [f"{_ANAL_PAY_MONTH_RU[m]} {year}, ₽" for m in _ANAL_PAY_MONTHS]
    )

    last_days = {m: calendar.monthrange(year, m)[1] for m in _ANAL_PAY_MONTHS}
    max_day = max(last_days.values())

    for day in range(1, max_day + 1):
        row = [day]
        for m in _ANAL_PAY_MONTHS:
            if day <= last_days[m]:
                row.append(int(daily_by_month.get(m, {}).get(day, 0)))
            else:
                row.append(None)
        ws_data.append(row)

    totals = ["Итого"]
    for col_idx, m in enumerate(_ANAL_PAY_MONTHS, start=2):
        col_letter = get_column_letter(col_idx)
        totals.append(f"=SUM({col_letter}2:{col_letter}{last_days[m] + 1})")
    ws_data.append(totals)
    total_row = max_day + 2

    for col in range(1, 5):
        cell = ws_data.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for r in range(2, total_row + 1):
        for c in range(1, 5):
            cell = ws_data.cell(row=r, column=c)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if r == total_row:
                cell.fill = total_fill
                cell.font = Font(bold=True)
            if c > 1:
                cell.number_format = '#,##0'

    for col in ws_data.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_data.column_dimensions[col_letter].width = min(max(max_len + 2, 16), 50)

    ws_charts = wb.create_sheet("Графики", 0)
    ws_charts.sheet_view.showGridLines = False
    wb.active = ws_charts

    chart_fills = ("5B9BD5", "70AD47", "ED7D31")
    chart_row = 1
    for idx, month in enumerate(_ANAL_PAY_MONTHS):
        days = last_days[month]
        month_total = sum(daily_by_month.get(month, {}).get(d, 0) for d in range(1, days + 1))
        total_label = f"{month_total:,}".replace(",", " ")
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = f"{_ANAL_PAY_MONTH_RU[month]} {year} — {total_label} ₽"
        chart.y_axis.title = "Платежи, ₽"
        chart.x_axis.title = "День"
        chart.y_axis.numFmt = '#,##0'
        chart.style = 10
        chart.legend = None
        chart.width = 22
        chart.height = 10
        data = Reference(ws_data, min_col=idx + 2, min_row=1, max_row=days + 1)
        cats = Reference(ws_data, min_col=1, min_row=2, max_row=days + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        if chart.series:
            chart.series[0].graphicalProperties.solidFill = chart_fills[idx]
        ws_charts.add_chart(chart, f"A{chart_row}")
        chart_row += 20

    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _check_recurent_period(now: datetime | None = None) -> tuple[datetime, datetime, date]:
    """1 августа — сегодня; граница подсветки — 25 августа того же года, что и начало периода."""
    now = now or datetime.now()
    end_date = datetime(now.year, now.month, now.day, 23, 59, 59)
    if now.month >= 8:
        period_year = now.year
    else:
        period_year = now.year - 1
    start_date = datetime(period_year, 8, 1, 0, 0, 0)
    split_date = date(period_year, 8, 25)
    return start_date, end_date, split_date


async def _fetch_daily_revenue(start_date: datetime, end_date: datetime) -> dict[date, int]:
    """Сумма успешных платежей (₽) по дням из всех таблиц оплат."""
    daily: dict[date, int] = {}
    current = start_date.date()
    while current <= end_date.date():
        daily[current] = 0
        current += timedelta(days=1)

    def _add(dt: datetime | None, rub) -> None:
        if dt is None or rub is None:
            return
        day = dt.date()
        if day not in daily:
            return
        daily[day] += int(rub)

    async with AsyncSessionLocal() as session:
        rub_models = (
            Payments,
            PaymentsCards,
            PaymentsPlategaCrypto,
            PaymentsWataSBP,
            PaymentsWataCard,
            PaymentsFkSBP,
        )
        for model in rub_models:
            stmt = select(model.time_created, model.amount).where(
                model.status.in_(_PAYMENT_OK_STATUSES),
                model.time_created.between(start_date, end_date),
            )
            for tc, amt in (await session.execute(stmt)).all():
                _add(tc, amt)

        stmt_stars = select(PaymentsStars.time_created, PaymentsStars.amount).where(
            PaymentsStars.status.in_(_PAYMENT_OK_STATUSES),
            PaymentsStars.time_created.between(start_date, end_date),
        )
        for tc, amt in (await session.execute(stmt_stars)).all():
            _add(tc, _stars_amount_to_rub(amt))

        stmt_crypto = select(
            PaymentsCryptobot.time_created,
            PaymentsCryptobot.amount,
        ).where(
            PaymentsCryptobot.status.in_(_PAYMENT_OK_STATUSES),
            PaymentsCryptobot.time_created.between(start_date, end_date),
        )
        for tc, amt in (await session.execute(stmt_crypto)).all():
            if amt is not None and float(amt) > 0.02:
                _add(tc, int(round(float(amt))))

        stmt_rec = select(PlategaRecurent.time_created, PlategaRecurent.amount).where(
            func.lower(PlategaRecurent.status).in_(_PAYMENT_OK_STATUSES),
            PlategaRecurent.time_created.between(start_date, end_date),
        )
        for tc, amt in (await session.execute(stmt_rec)).all():
            _add(tc, amt)

    return daily


def _sync_build_check_recurent_excel(
    daily: dict[date, int],
    split_date: date,
    start_date: datetime,
    end_date: datetime,
) -> str:
    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Данные"

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    ws_data.append(["Дата", "Выручка, ₽"])
    sorted_days = sorted(daily.keys())
    for day in sorted_days:
        ws_data.append([day.strftime("%d.%m.%Y"), int(daily[day])])

    total_row = len(sorted_days) + 2
    ws_data.append(["Итого", f"=SUM(B2:B{total_row - 1})"])

    for col in range(1, 3):
        cell = ws_data.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for r in range(2, total_row + 1):
        row_fill = None
        if r < total_row:
            row_date = sorted_days[r - 2]
            row_fill = green_fill if row_date >= split_date else yellow_fill
        else:
            row_fill = total_fill

        for c in range(1, 3):
            cell = ws_data.cell(row=r, column=c)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if row_fill:
                cell.fill = row_fill
            if r == total_row:
                cell.font = Font(bold=True)
            if c == 2 and r < total_row:
                cell.number_format = '#,##0'

    for col in ws_data.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_data.column_dimensions[col_letter].width = min(max(max_len + 2, 14), 50)

    ws_data.freeze_panes = "A2"

    ws_chart = wb.create_sheet("График", 0)
    ws_chart.sheet_view.showGridLines = False
    wb.active = ws_chart

    if sorted_days:
        total_revenue = sum(daily.values())
        total_label = f"{total_revenue:,}".replace(",", " ")
        period_label = (
            f"{start_date.strftime('%d.%m.%Y')} — {end_date.strftime('%d.%m.%Y')}"
        )
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = f"Выручка по дням ({period_label}) — {total_label} ₽"
        chart.y_axis.title = "Выручка, ₽"
        chart.x_axis.title = "Дата"
        chart.y_axis.numFmt = '#,##0'
        chart.style = 10
        chart.legend = None
        chart.width = max(18, min(len(sorted_days) * 0.45, 60))
        chart.height = 12

        data_ref = Reference(ws_data, min_col=2, min_row=1, max_row=len(sorted_days) + 1)
        cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=len(sorted_days) + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        if chart.series:
            series = chart.series[0]
            data_points = []
            for i, day in enumerate(sorted_days):
                pt = DataPoint(idx=i)
                pt.graphicalProperties.solidFill = (
                    "CCFFCC" if day >= split_date else "FFFF00"
                )
                data_points.append(pt)
            series.dPt = data_points

        ws_chart.add_chart(chart, "A1")

    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


@router.message(Command(commands=['stat']))
async def stat_command(message: Message):
    """Статистика по пользователям с указанным Ref или stamp (админы и CHECKER_IDS)."""
    if message.from_user.id not in ADMIN_IDS | CHECKER_IDS:
        return

    args = message.text.split()
    if len(args) < 2:
        await message.answer("❌ Использование: /stat <аргумент>")
        return

    arg = args[1].strip()
    total, with_sub, with_tarif, with_tarif_not_blocked, total_payments, source = await sql.get_stat_by_ref_or_stamp(arg)

    if total is None:
        await message.answer(f"{arg} - нет совпадений")
    else:
        await message.answer(
            f"{arg} {total} {with_sub} {with_tarif} {with_tarif_not_blocked} - {total_payments} руб"
        )


@router.message(Command(commands=['anal_export']))
async def analytics_export(message: Message):
    if message.from_user.id not in ADMIN_IDS:
        return

    await message.answer("🔄 Формирую помесячную аналитику...")

    try:
        now = datetime.now()
        current_year = now.year
        current_month = now.month

        months = [(current_year, month) for month in range(1, current_month + 1)]

        monthly_data = {}
        daily_data_by_month = {}
        cumulative_revenue = 0

        async with AsyncSessionLocal() as session:
            payment_amounts = await _discover_payment_amounts(session)
            await session.commit()

        for year, month in months:
            start_date = datetime(year, month, 1, 0, 0, 0)
            last_day = calendar.monthrange(year, month)[1]
            end_date = datetime(year, month, last_day, 23, 59, 59)
            month_key = start_date.strftime('%B %Y')

            async with AsyncSessionLocal() as session:
                # --- Новые пользователи за месяц ---
                stmt_new_users = select(Users).where(
                    Users.create_user.between(start_date, end_date)
                )
                result = await session.execute(stmt_new_users)
                new_users = result.scalars().all()

                new_total = []
                new_zaliv = []
                new_saraf = []
                key_total = []
                key_zaliv = []
                key_saraf = []
                connect_total = []
                connect_zaliv = []
                connect_saraf = []
                set_new_total = set()
                set_new_zaliv = set()
                set_new_saraf = set()

                daily_stats = {day: {'new': 0, 'key': 0, 'connect': 0, 'paid': 0} for day in range(1, last_day + 1)}

                for user in new_users:
                    is_zaliv = (user.stamp != '')
                    uid = user.user_id
                    create_day = user.create_user.day

                    new_total.append(uid)
                    set_new_total.add(uid)
                    if is_zaliv:
                        new_zaliv.append(uid)
                        set_new_zaliv.add(uid)
                    else:
                        new_saraf.append(uid)
                        set_new_saraf.add(uid)

                    if user.in_panel:
                        key_total.append(uid)
                        if is_zaliv:
                            key_zaliv.append(uid)
                        else:
                            key_saraf.append(uid)

                    if user.is_connect:
                        connect_total.append(uid)
                        if is_zaliv:
                            connect_zaliv.append(uid)
                        else:
                            connect_saraf.append(uid)

                    daily_stats[create_day]['new'] += 1
                    if user.in_panel:
                        daily_stats[create_day]['key'] += 1
                    if user.is_connect:
                        daily_stats[create_day]['connect'] += 1

                # --- Множество плативших ---
                stmt_paid_main = select(Payments.user_id).distinct().where(
                    Payments.status == 'confirmed',
                    Payments.amount != 1
                )
                paid_main = {row[0] for row in (await session.execute(stmt_paid_main)).all()}

                stmt_paid_stars = select(PaymentsStars.user_id).distinct().where(
                    PaymentsStars.status == 'confirmed'
                )
                paid_stars = {row[0] for row in (await session.execute(stmt_paid_stars)).all()}

                stmt_paid_crypto = select(PaymentsCryptobot.user_id).distinct().where(
                    PaymentsCryptobot.status == 'paid',
                    PaymentsCryptobot.amount > 0.02
                )
                paid_crypto = {row[0] for row in (await session.execute(stmt_paid_crypto)).all()}

                stmt_paid_cards = select(PaymentsCards.user_id).distinct().where(
                    PaymentsCards.status == 'confirmed',
                    PaymentsCards.amount != 1
                )
                paid_cards = {row[0] for row in (await session.execute(stmt_paid_cards)).all()}

                stmt_paid_platega_crypto = select(PaymentsPlategaCrypto.user_id).distinct().where(
                    PaymentsPlategaCrypto.status == 'confirmed',
                    PaymentsPlategaCrypto.amount != 1  # если нужно исключить тестовые платежи
                )
                paid_platega_crypto = {row[0] for row in (await session.execute(stmt_paid_platega_crypto)).all()}

                stmt_paid_wata_sbp = select(PaymentsWataSBP.user_id).distinct().where(
                    PaymentsWataSBP.status == 'confirmed',
                    PaymentsWataSBP.amount != 1,
                )
                paid_wata_sbp = {row[0] for row in (await session.execute(stmt_paid_wata_sbp)).all()}

                stmt_paid_wata_card = select(PaymentsWataCard.user_id).distinct().where(
                    PaymentsWataCard.status == 'confirmed',
                    PaymentsWataCard.amount != 1,
                )
                paid_wata_card = {row[0] for row in (await session.execute(stmt_paid_wata_card)).all()}

                stmt_paid_fk = select(PaymentsFkSBP.user_id).distinct().where(
                    PaymentsFkSBP.status == 'confirmed',
                    PaymentsFkSBP.amount != 1,
                )
                paid_fk = {row[0] for row in (await session.execute(stmt_paid_fk)).all()}

                all_paid_users = paid_main.union(paid_stars).union(paid_crypto).union(paid_cards).union(
                    paid_platega_crypto).union(paid_wata_sbp).union(paid_wata_card).union(paid_fk)

                for uid in set_new_total:
                    if uid in all_paid_users:
                        # найдём день регистрации
                        for user in new_users:
                            if user.user_id == uid:
                                daily_stats[user.create_user.day]['paid'] += 1
                                break

                # --- Платежи новых пользователей за этот месяц ---
                new_payments_amounts = []

                # Основные
                stmt_main_new = select(Payments.user_id, Payments.amount).where(
                    Payments.time_created.between(start_date, end_date),
                    Payments.amount != 1,
                    Payments.status == 'confirmed'
                )
                for uid, amt in (await session.execute(stmt_main_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                # Звёзды
                stmt_stars_new = select(PaymentsStars.user_id, PaymentsStars.amount).where(
                    PaymentsStars.time_created.between(start_date, end_date),
                    PaymentsStars.status == 'confirmed'
                )
                for uid, amt in (await session.execute(stmt_stars_new)).all():
                    if uid in set_new_total:
                        rub = convert_stars_to_rub(amt)
                        if rub:
                            new_payments_amounts.append((uid, rub))

                # Cryptobot (amount в БД уже в рублях)
                stmt_crypto_new = select(
                    PaymentsCryptobot.user_id,
                    PaymentsCryptobot.amount,
                ).where(
                    PaymentsCryptobot.time_created.between(start_date, end_date),
                    PaymentsCryptobot.status == 'paid',
                    PaymentsCryptobot.amount > 0.02
                )
                for uid, amt in (await session.execute(stmt_crypto_new)).all():
                    if uid in set_new_total:
                        rub = int(round(amt))
                        if rub:
                            new_payments_amounts.append((uid, rub))

                stmt_cards_new = select(PaymentsCards.user_id, PaymentsCards.amount).where(
                    PaymentsCards.time_created.between(start_date, end_date),
                    PaymentsCards.amount != 1,
                    PaymentsCards.status == 'confirmed'
                )
                for uid, amt in (await session.execute(stmt_cards_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                # Platega Crypto (новые пользователи)
                stmt_platega_crypto_new = select(PaymentsPlategaCrypto.user_id,
                                                 PaymentsPlategaCrypto.amount).where(
                    PaymentsPlategaCrypto.time_created.between(start_date, end_date),
                    PaymentsPlategaCrypto.amount != 1,
                    PaymentsPlategaCrypto.status == 'confirmed'
                )
                for uid, amt in (await session.execute(stmt_platega_crypto_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                stmt_wata_sbp_new = select(PaymentsWataSBP.user_id, PaymentsWataSBP.amount).where(
                    PaymentsWataSBP.time_created.between(start_date, end_date),
                    PaymentsWataSBP.amount != 1,
                    PaymentsWataSBP.status == 'confirmed',
                )
                for uid, amt in (await session.execute(stmt_wata_sbp_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                stmt_wata_card_new = select(PaymentsWataCard.user_id, PaymentsWataCard.amount).where(
                    PaymentsWataCard.time_created.between(start_date, end_date),
                    PaymentsWataCard.amount != 1,
                    PaymentsWataCard.status == 'confirmed',
                )
                for uid, amt in (await session.execute(stmt_wata_card_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                stmt_fk_new = select(PaymentsFkSBP.user_id, PaymentsFkSBP.amount).where(
                    PaymentsFkSBP.time_created.between(start_date, end_date),
                    PaymentsFkSBP.amount != 1,
                    PaymentsFkSBP.status == 'confirmed',
                )
                for uid, amt in (await session.execute(stmt_fk_new)).all():
                    if uid in set_new_total:
                        new_payments_amounts.append((uid, amt))

                pay_sum_total = 0
                pay_sum_zaliv = 0
                pay_sum_saraf = 0
                pay_users_total = set()
                pay_users_zaliv = set()
                pay_users_saraf = set()

                for uid, amount in new_payments_amounts:
                    pay_sum_total += amount
                    pay_users_total.add(uid)
                    if uid in set_new_zaliv:
                        pay_sum_zaliv += amount
                        pay_users_zaliv.add(uid)
                    elif uid in set_new_saraf:
                        pay_sum_saraf += amount
                        pay_users_saraf.add(uid)

                # --- Общие платежи за месяц (все пользователи) ---
                all_payments = []  # (amount, is_gift)

                # Основные
                stmt_main_all = select(Payments.amount, Payments.is_gift).where(
                    Payments.time_created.between(start_date, end_date),
                    Payments.amount != 1,
                    Payments.status == 'confirmed'
                )
                for amount, is_gift in (await session.execute(stmt_main_all)).all():
                    all_payments.append((amount, is_gift))

                # Звёзды
                stmt_stars_all = select(PaymentsStars.amount, PaymentsStars.is_gift).where(
                    PaymentsStars.time_created.between(start_date, end_date),
                    PaymentsStars.status == 'confirmed'
                )
                for amount, is_gift in (await session.execute(stmt_stars_all)).all():
                    rub = convert_stars_to_rub(amount)
                    if rub:
                        all_payments.append((rub, is_gift))

                # Cryptobot (amount в БД уже в рублях)
                stmt_crypto_all = select(
                    PaymentsCryptobot.amount,
                    PaymentsCryptobot.is_gift
                ).where(
                    PaymentsCryptobot.time_created.between(start_date, end_date),
                    PaymentsCryptobot.status == 'paid',
                    PaymentsCryptobot.amount > 0.02
                )
                for amount, is_gift in (await session.execute(stmt_crypto_all)).all():
                    rub = int(round(amount))
                    if rub:
                        all_payments.append((rub, is_gift))

                stmt_cards_all = select(PaymentsCards.amount, PaymentsCards.is_gift).where(
                    PaymentsCards.time_created.between(start_date, end_date),
                    PaymentsCards.amount != 1,
                    PaymentsCards.status == 'confirmed'
                )
                for amount, is_gift in (await session.execute(stmt_cards_all)).all():
                    all_payments.append((amount, is_gift))

                # Platega Crypto (все пользователи)
                stmt_platega_crypto_all = select(PaymentsPlategaCrypto.amount, PaymentsPlategaCrypto.is_gift).where(
                    PaymentsPlategaCrypto.time_created.between(start_date, end_date),
                    PaymentsPlategaCrypto.amount != 1,
                    PaymentsPlategaCrypto.status == 'confirmed'
                )
                for amount, is_gift in (await session.execute(stmt_platega_crypto_all)).all():
                    all_payments.append((amount, is_gift))

                stmt_wata_sbp_all = select(PaymentsWataSBP.amount, PaymentsWataSBP.is_gift).where(
                    PaymentsWataSBP.time_created.between(start_date, end_date),
                    PaymentsWataSBP.amount != 1,
                    PaymentsWataSBP.status == 'confirmed',
                )
                for amount, is_gift in (await session.execute(stmt_wata_sbp_all)).all():
                    all_payments.append((amount, is_gift))

                stmt_wata_card_all = select(PaymentsWataCard.amount, PaymentsWataCard.is_gift).where(
                    PaymentsWataCard.time_created.between(start_date, end_date),
                    PaymentsWataCard.amount != 1,
                    PaymentsWataCard.status == 'confirmed',
                )
                for amount, is_gift in (await session.execute(stmt_wata_card_all)).all():
                    all_payments.append((amount, is_gift))

                stmt_fk_all = select(PaymentsFkSBP.amount, PaymentsFkSBP.is_gift).where(
                    PaymentsFkSBP.time_created.between(start_date, end_date),
                    PaymentsFkSBP.amount != 1,
                    PaymentsFkSBP.status == 'confirmed',
                )
                for amount, is_gift in (await session.execute(stmt_fk_all)).all():
                    all_payments.append((amount, is_gift))

                total_revenue = sum(p[0] for p in all_payments)
                total_payments_count = len(all_payments)
                aov = total_revenue / total_payments_count if total_payments_count else 0

                stmt_cumulative = select(func.count(Users.id)).where(
                    Users.create_user <= end_date
                )
                cumulative_users = (await session.execute(stmt_cumulative)).scalar() or 1
                # ARPU по нарастающей: выручка с января до конца месяца /
                # пользователи, созданные до конца этого месяца.
                cumulative_revenue += total_revenue
                arpu = cumulative_revenue / cumulative_users if cumulative_users else 0

                # Разбивка по суммам (динамически по всем номиналам из БД)
                amount_stats = {amt: {'count': 0, 'amount': 0} for amt in payment_amounts}
                gift_count = gift_amount = 0

                for amount, is_gift in all_payments:
                    if is_gift:
                        gift_count += 1
                        gift_amount += amount
                    else:
                        rub = _normalize_payment_amount(amount)
                        if rub is not None and rub in amount_stats:
                            amount_stats[rub]['count'] += 1
                            amount_stats[rub]['amount'] += rub

                month_row = {
                    'new_total': len(new_total),
                    'new_zaliv': len(new_zaliv),
                    'new_saraf': len(new_saraf),
                    'key_total': len(key_total),
                    'key_zaliv': len(key_zaliv),
                    'key_saraf': len(key_saraf),
                    'connect_total': len(connect_total),
                    'connect_zaliv': len(connect_zaliv),
                    'connect_saraf': len(connect_saraf),
                    'pay_new_sum_total': pay_sum_total,
                    'pay_new_users_total': len(pay_users_total),
                    'pay_new_sum_zaliv': pay_sum_zaliv,
                    'pay_new_users_zaliv': len(pay_users_zaliv),
                    'pay_new_sum_saraf': pay_sum_saraf,
                    'pay_new_users_saraf': len(pay_users_saraf),
                    'total_revenue': total_revenue,
                    'total_payments': total_payments_count,
                    'aov': aov,
                    'arpu': arpu,
                    'cumulative_users': cumulative_users,
                    'gift_count': gift_count,
                    'gift_amount': gift_amount,
                }
                for amt in payment_amounts:
                    month_row[f'sum_{amt}_count'] = amount_stats[amt]['count']
                    month_row[f'sum_{amt}_amount'] = amount_stats[amt]['amount']
                monthly_data[month_key] = month_row

                # --- Поденные данные (кумулятивные) ---
                stmt_before = select(Users.user_id, Users.in_panel, Users.is_connect).where(
                    Users.create_user < start_date
                )
                users_before = (await session.execute(stmt_before)).all()
                cum_users_before = len(users_before)
                cum_key_before = sum(1 for u in users_before if u.in_panel)
                cum_connect_before = sum(1 for u in users_before if u.is_connect)

                # Закрываем read-транзакцию до тяжёлой обработки в Python (меньше времени удержания lock)
                await session.commit()

                daily_cumulative = []
                cum_users = cum_users_before
                cum_key = cum_key_before
                cum_connect = cum_connect_before

                for day in range(1, last_day + 1):
                    cum_users += daily_stats[day]['new']
                    cum_key += daily_stats[day]['key']
                    cum_connect += daily_stats[day]['connect']
                    daily_cumulative.append({
                        'day': day,
                        'cum_users': cum_users,
                        'cum_key': cum_key,
                        'cum_connect': cum_connect,
                        'new': daily_stats[day]['new'],
                        'key': daily_stats[day]['key'],
                        'connect': daily_stats[day]['connect'],
                        'paid': daily_stats[day]['paid']
                    })

                daily_data_by_month[month_key] = daily_cumulative

        export_path = await asyncio.to_thread(
            _sync_build_analytics_excel, monthly_data, daily_data_by_month, payment_amounts
        )
        try:
            await message.answer_document(
                document=FSInputFile(
                    export_path,
                    filename=f"analytics_{current_year}_{current_month}.xlsx",
                ),
                caption=f"📊 Помесячная аналитика с января {current_year} по {now.strftime('%B %Y')}",
            )
        finally:
            try:
                os.remove(export_path)
            except OSError:
                pass

        logger.info(f"Админ {message.from_user.id} выгрузил помесячную аналитику")

    except Exception as e:
        logger.exception("Ошибка при экспорте помесячной аналитики")
        await message.answer(f"❌ Ошибка: {str(e)}")


@router.message(Command(commands=['anal_payment']))
async def anal_payment_command(message: Message):
    """Excel: графики успешных платежей (₽) по дням за июнь, июль, август."""
    if message.from_user.id not in ADMIN_IDS:
        await message.answer("❌ Команда доступна только администраторам.")
        return

    await message.answer("🔄 Формирую графики платежей за июнь–август...")

    try:
        now = datetime.now()
        year = now.year if now.month >= 6 else now.year - 1
        start_date = datetime(year, 6, 1, 0, 0, 0)
        end_date = datetime(year, 8, 31, 23, 59, 59)

        last_days = {m: calendar.monthrange(year, m)[1] for m in _ANAL_PAY_MONTHS}
        daily_by_month = {
            m: {d: 0 for d in range(1, last_days[m] + 1)}
            for m in _ANAL_PAY_MONTHS
        }

        def _add(dt, rub):
            if dt is None or rub is None:
                return
            if dt.year != year or dt.month not in daily_by_month:
                return
            day = dt.day
            if day in daily_by_month[dt.month]:
                daily_by_month[dt.month][day] += int(rub)

        async with AsyncSessionLocal() as session:
            rub_models = (
                Payments,
                PaymentsCards,
                PaymentsPlategaCrypto,
                PaymentsWataSBP,
                PaymentsWataCard,
                PaymentsFkSBP,
            )
            for model in rub_models:
                stmt = select(model.time_created, model.amount).where(
                    model.status.in_(_PAYMENT_OK_STATUSES),
                    model.time_created.between(start_date, end_date),
                )
                for tc, amt in (await session.execute(stmt)).all():
                    _add(tc, amt)

            stmt_stars = select(PaymentsStars.time_created, PaymentsStars.amount).where(
                PaymentsStars.status.in_(_PAYMENT_OK_STATUSES),
                PaymentsStars.time_created.between(start_date, end_date),
            )
            for tc, amt in (await session.execute(stmt_stars)).all():
                _add(tc, _stars_amount_to_rub(amt))

            stmt_crypto = select(
                PaymentsCryptobot.time_created,
                PaymentsCryptobot.amount,
            ).where(
                PaymentsCryptobot.status.in_(_PAYMENT_OK_STATUSES),
                PaymentsCryptobot.time_created.between(start_date, end_date),
            )
            for tc, amt in (await session.execute(stmt_crypto)).all():
                if amt is not None and float(amt) > 0.02:
                    _add(tc, int(round(float(amt))))

            stmt_rec = select(PlategaRecurent.time_created, PlategaRecurent.amount).where(
                func.lower(PlategaRecurent.status).in_(_PAYMENT_OK_STATUSES),
                PlategaRecurent.time_created.between(start_date, end_date),
            )
            for tc, amt in (await session.execute(stmt_rec)).all():
                _add(tc, amt)

        export_path = await asyncio.to_thread(
            _sync_build_anal_payment_excel, year, daily_by_month
        )
        june_sum = sum(daily_by_month[6].values())
        july_sum = sum(daily_by_month[7].values())
        aug_sum = sum(daily_by_month[8].values())

        def _fmt_rub(n: int) -> str:
            return f"{n:,}".replace(",", " ")

        try:
            await message.answer_document(
                document=FSInputFile(
                    export_path,
                    filename=f"anal_payment_{year}_06-08.xlsx",
                ),
                caption=(
                    f"📊 Платежи по дням, {year}\n"
                    f"Июнь: {_fmt_rub(june_sum)} ₽\n"
                    f"Июль: {_fmt_rub(july_sum)} ₽\n"
                    f"Август: {_fmt_rub(aug_sum)} ₽"
                ),
            )
        finally:
            try:
                os.remove(export_path)
            except OSError:
                pass

        logger.info(f"Админ {message.from_user.id} выгрузил графики платежей /anal_payment")

    except Exception as e:
        logger.exception("Ошибка при экспорте графиков платежей")
        await message.answer(f"❌ Ошибка: {str(e)}")


@router.message(Command(commands=['check_recurent']))
async def check_recurent_command(message: Message):
    """Excel: выручка по дням с 1 августа; до 25.08 — жёлтый, с 25.08 — зелёный + график."""
    if message.from_user.id not in ADMIN_IDS:
        await message.answer("❌ Команда доступна только администраторам.")
        return

    await message.answer("🔄 Формирую отчёт по выручке с 1 августа...")

    try:
        now = datetime.now()
        start_date, end_date, split_date = _check_recurent_period(now)
        daily = await _fetch_daily_revenue(start_date, end_date)

        export_path = await asyncio.to_thread(
            _sync_build_check_recurent_excel,
            daily,
            split_date,
            start_date,
            end_date,
        )

        total = sum(daily.values())
        before_split = sum(v for d, v in daily.items() if d < split_date)
        from_split = sum(v for d, v in daily.items() if d >= split_date)

        def _fmt_rub(n: int) -> str:
            return f"{n:,}".replace(",", " ")

        try:
            await message.answer_document(
                document=FSInputFile(
                    export_path,
                    filename=(
                        f"check_recurent_{start_date.strftime('%d.%m.%y')}_"
                        f"{end_date.strftime('%d.%m.%y')}.xlsx"
                    ),
                ),
                caption=(
                    f"📊 Выручка по дням\n"
                    f"Период: {start_date.strftime('%d.%m.%Y')} — {end_date.strftime('%d.%m.%Y')}\n"
                    f"🟡 До {split_date.strftime('%d.%m')}: {_fmt_rub(before_split)} ₽\n"
                    f"🟢 С {split_date.strftime('%d.%m')}: {_fmt_rub(from_split)} ₽\n"
                    f"💰 Итого: {_fmt_rub(total)} ₽"
                ),
            )
        finally:
            try:
                os.remove(export_path)
            except OSError:
                pass

        logger.info(f"Админ {message.from_user.id} выгрузил /check_recurent")

    except Exception as e:
        logger.exception("Ошибка при экспорте /check_recurent")
        await message.answer(f"❌ Ошибка: {str(e)}")
