from datetime import date, datetime, timedelta, timezone
from collections import defaultdict
from typing import Any, List, Optional, Tuple
from zoneinfo import ZoneInfo
import asyncio
import os
import tempfile

import openpyxl
from aiogram import F, Router
from openpyxl.styles import Alignment, Border, Side, PatternFill

from bot import bot, sql, x3
from config import ADMIN_IDS, MIRROR_SUB_LINK, TRUE_SUB_LINK
from config_bd.models import Users
from config_bd.utils import (
    _billing_duration_from_amount_fallback,
    _parse_traffic_duration,
    _payload_duration_to_panel_days,
)
from keyboard import STYLE_DANGER, STYLE_PRIMARY, STYLE_SUCCESS
from logging_config import logger
from aiogram.types import (
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
)
from aiogram.filters import Command
from telegram_ids import is_telegram_chat_id
from wl_traffic.service import (
    fetch_panel_user,
    is_forever_duration,
    is_forever_end_date,
    panel_username_for_billing_uid,
    reassign_to_active_squad,
    user_on_active_squad,
)

router = Router()

_EXCEL_COL_WIDTH_MAX = 255

_USERS_EXPORT_COLUMNS_DEFAULT = (
    'id', 'user_id', 'ref', 'is_delete', 'in_panel', 'is_connect',
    'create_user', 'reserve_field', 'subscription_end_date',
    'last_notification_date',
    'last_broadcast_date', 'stamp', 'ttclid',
    'field_bool_3',
)


def _user_sheet_column_names(users_full_columns: bool) -> list[str]:
    if users_full_columns:
        return [c.key for c in Users.__table__.columns]
    return list(_USERS_EXPORT_COLUMNS_DEFAULT)


def _excel_scalar(value):
    if value is None:
        return value
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return value


async def _export_database_to_excel_impl(message: Message, *, users_full_columns: bool) -> None:
    """Экспорт базы в Excel; при users_full_columns на листе users все колонки таблицы."""
    if message.from_user.id not in ADMIN_IDS:
        await message.answer("❌ Эта команда доступна только администраторам.")
        return

    try:
        start_msg = (
            "🔄 Начинаю экспорт базы данных (лист users — все колонки)..."
            if users_full_columns
            else "🔄 Начинаю экспорт базы данных..."
        )
        await message.answer(start_msg)

        snapshot = await sql.get_export_snapshot()

        def _sync_build_export() -> str:
            users_list = snapshot["users"]
            payments_list = snapshot["payments"]
            payments_cards_list = snapshot["payments_cards"]
            payments_platega_crypto_list = snapshot["payments_platega_crypto"]
            payments_fk_sbp_list = snapshot["payments_fk_sbp"]
            payments_wata_sbp_list = snapshot["payments_wata_sbp"]
            payments_wata_card_list = snapshot["payments_wata_card"]
            payments_stars_list = snapshot["payments_stars"]
            payments_cryptobot_list = snapshot["payments_cryptobot"]
            gifts_list = snapshot["gifts"]
            online_list = snapshot["online"]
            white_counter_list = snapshot["white_counter"]
            platega_autopay_list = snapshot["platega_autopay_subscriptions"]
            platega_recurent_list = snapshot["platega_recurent"]

            wb = openpyxl.Workbook()
            if 'Sheet' in wb.sheetnames:
                wb.remove(wb['Sheet'])
            
            # --- Лист USERS ---
            ws_users = wb.create_sheet(title="users")
            users_columns = _user_sheet_column_names(users_full_columns)
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))

            for col_num, title in enumerate(users_columns, 1):
                cell = ws_users.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border

            for row_num, user in enumerate(users_list, 2):
                row_data = [_excel_scalar(getattr(user, name)) for name in users_columns]
                for col_num, value in enumerate(row_data, 1):
                    cell = ws_users.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border

            for col in ws_users.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_users.column_dimensions[col_letter].width = min(max_len + 2, _EXCEL_COL_WIDTH_MAX)
            
            # --- Лист PAYMENTS (Platega) ---
            ws_payments = wb.create_sheet(title="payments_sbp")
            payments_columns = ['ID', 'User ID', 'Amount', 'Time Created', 'Is Gift', 'Status', 'Transaction_Id']
            for col_num, title in enumerate(payments_columns, 1):
                cell = ws_payments.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, pay in enumerate(payments_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_payments.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_payments.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_payments.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            
            
            # --- Лист PAYMENTS_CARDS (платежи по картам) ---
            ws_payments_cards = wb.create_sheet(title="payments_cards")
            cards_columns = ['ID', 'User ID', 'Amount', 'Time Created', 'Is Gift', 'Status', 'Transaction_Id',
                             'Payload']
            for col_num, title in enumerate(cards_columns, 1):
                cell = ws_payments_cards.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, pay in enumerate(payments_cards_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id, pay.payload
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_payments_cards.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_payments_cards.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_payments_cards.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            
            # --- Лист PAYMENTS_STARS ---
            ws_payments_stars = wb.create_sheet(title="payments_stars")
            stars_columns = ['ID', 'User ID', 'Amount (Stars)', 'Time Created', 'Is Gift', 'Status']
            for col_num, title in enumerate(stars_columns, 1):
                cell = ws_payments_stars.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, ps in enumerate(payments_stars_list, 2):
                row_data = [
                    ps.id, ps.user_id, ps.amount, ps.time_created,
                    ps.is_gift, ps.status
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_payments_stars.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_payments_stars.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_payments_stars.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            
            
            # --- Лист PAYMENTS_PLATEGA_CRYPTO ---
            ws_platega_crypto = wb.create_sheet(title="payments_platega_crypto")
            platega_crypto_columns = ['ID', 'User ID', 'Amount', 'Time Created', 'Is Gift', 'Status', 'Transaction_Id',
                                      'Payload']
            for col_num, title in enumerate(platega_crypto_columns, 1):
                cell = ws_platega_crypto.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, pay in enumerate(payments_platega_crypto_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id, pay.payload
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_platega_crypto.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_platega_crypto.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_platega_crypto.column_dimensions[col_letter].width = min(max_len + 2, 50)

            ws_fk_sbp = wb.create_sheet(title="payments_fk_sbp")
            fk_columns = [
                'ID', 'User ID', 'Amount', 'Time Created', 'Is Gift', 'Status',
                'Transaction_Id', 'FK_Order_Id', 'Nonce', 'Signature', 'Method', 'Payload',
            ]
            for col_num, title in enumerate(fk_columns, 1):
                cell = ws_fk_sbp.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border

            for row_num, pay in enumerate(payments_fk_sbp_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id, pay.fk_order_id,
                    pay.nonce, pay.signature, pay.method, pay.payload,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_fk_sbp.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border

            for col in ws_fk_sbp.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_fk_sbp.column_dimensions[col_letter].width = min(max_len + 2, 50)

            wata_columns = ['ID', 'User ID', 'Amount', 'Time Created', 'Is Gift', 'Status', 'Transaction_Id', 'Payload']
            ws_wata_sbp = wb.create_sheet(title="payments_wata_sbp")
            for col_num, title in enumerate(wata_columns, 1):
                cell = ws_wata_sbp.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            for row_num, pay in enumerate(payments_wata_sbp_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id, pay.payload,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_wata_sbp.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            for col in ws_wata_sbp.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_wata_sbp.column_dimensions[col_letter].width = min(max_len + 2, 50)

            ws_wata_card = wb.create_sheet(title="payments_wata_card")
            for col_num, title in enumerate(wata_columns, 1):
                cell = ws_wata_card.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            for row_num, pay in enumerate(payments_wata_card_list, 2):
                row_data = [
                    pay.id, pay.user_id, pay.amount, pay.time_created,
                    pay.is_gift, pay.status, pay.transaction_id, pay.payload,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 4 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_wata_card.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            for col in ws_wata_card.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_wata_card.column_dimensions[col_letter].width = min(max_len + 2, 50)

            # --- Лист PAYMENTS_CRYPTOBOT ---
            ws_payments_cryptobot = wb.create_sheet(title="payments_cryptobot")
            crypto_columns = [
                'ID', 'User ID', 'Amount', 'Currency', 'Time Created',
                'Is Gift', 'Status', 'Invoice ID', 'Payload'
            ]
            for col_num, title in enumerate(crypto_columns, 1):
                cell = ws_payments_cryptobot.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, pc in enumerate(payments_cryptobot_list, 2):
                row_data = [
                    pc.id, pc.user_id, pc.amount, pc.currency, pc.time_created,
                    pc.is_gift, pc.status, pc.invoice_id, pc.payload
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 5 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_payments_cryptobot.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_payments_cryptobot.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_payments_cryptobot.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            # --- Лист GIFTS ---
            ws_gifts = wb.create_sheet(title="gifts")
            gifts_columns = ['gift_id', 'giver_id', 'duration', 'recepient_id', 'white_flag', 'flag']
            for col_num, title in enumerate(gifts_columns, 1):
                cell = ws_gifts.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, gift in enumerate(gifts_list, 2):
                row_data = [
                    gift.gift_id, gift.giver_id, gift.duration,
                    gift.recepient_id, gift.white_flag, gift.flag
                ]
                for col_num, value in enumerate(row_data, 1):
                    cell = ws_gifts.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_gifts.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_gifts.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            # --- Лист ONLINE ---
            ws_online = wb.create_sheet(title="online")
            online_columns = [
                'ID', 'Дата сбора', 'Всего в панели', 'Активны сегодня',
                'Платных', 'Триальных', 'С активной подпиской',
            ]
            for col_num, title in enumerate(online_columns, 1):
                cell = ws_online.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, rec in enumerate(online_list, 2):
                row_data = [
                    rec.online_id, rec.online_date, rec.users_panel,
                    rec.users_active, rec.users_pay, rec.users_trial,
                    rec.users_subscribed,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 2 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_online.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_online.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_online.column_dimensions[col_letter].width = min(max_len + 2, 50)
            
            # --- Лист WHITE_COUNTER ---
            ws_white_counter = wb.create_sheet(title="white_counter")
            wc_columns = ['ID', 'User ID', 'Time Created']
            for col_num, title in enumerate(wc_columns, 1):
                cell = ws_white_counter.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border
            
            for row_num, wc in enumerate(white_counter_list, 2):
                row_data = [wc.id, wc.user_id, wc.time_created]
                for col_num, value in enumerate(row_data, 1):
                    if col_num == 3 and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_white_counter.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
            
            for col in ws_white_counter.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_white_counter.column_dimensions[col_letter].width = min(max_len + 2, 50)

            # --- Лист PLATEGA_AUTOPAY (рекуррентные подписки) ---
            ws_platega_autopay = wb.create_sheet(title="platega_autopay")
            autopay_columns = [
                'ID', 'User ID', 'Subscription ID', 'Duration', 'Amount', 'Status',
                'Next Charge At', 'Time Created', 'Time Cancelled', 'Cancel Reason',
                'White', 'Source', 'Payload',
            ]
            for col_num, title in enumerate(autopay_columns, 1):
                cell = ws_platega_autopay.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border

            for row_num, row in enumerate(platega_autopay_list, 2):
                row_data = [
                    row.id, row.user_id, row.subscription_id, row.duration, row.amount,
                    row.status, row.next_charge_at, row.time_created, row.time_cancelled,
                    row.cancel_reason, row.white, row.source, row.payload,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num in (7, 8, 9) and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_platega_autopay.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border

            for col in ws_platega_autopay.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_platega_autopay.column_dimensions[col_letter].width = min(max_len + 2, 50)

            # --- Лист PLATEGA_RECURENT (списания по рекуррентам) ---
            ws_platega_recurent = wb.create_sheet(title="platega_recurent")
            recurent_columns = [
                'ID', 'User ID', 'Subscription ID', 'Transaction ID', 'Amount', 'Currency',
                'Status', 'Payment Method', 'Next Charge At', 'Time Created', 'Processed', 'Payload',
            ]
            for col_num, title in enumerate(recurent_columns, 1):
                cell = ws_platega_recurent.cell(row=1, column=col_num, value=title)
                cell.alignment = header_alignment
                cell.border = thin_border

            for row_num, row in enumerate(platega_recurent_list, 2):
                row_data = [
                    row.id, row.user_id, row.subscription_id, row.transaction_id, row.amount,
                    row.currency, row.status, row.payment_method, row.next_charge_at,
                    row.time_created, row.processed, row.payload,
                ]
                for col_num, value in enumerate(row_data, 1):
                    if col_num in (9, 10) and value and isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    cell = ws_platega_recurent.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border

            for col in ws_platega_recurent.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_platega_recurent.column_dimensions[col_letter].width = min(max_len + 2, 50)

            # Заморозка заголовков
            for ws in [ws_users, ws_payments, ws_payments_cards, ws_payments_stars, ws_platega_crypto,
                       ws_fk_sbp, ws_wata_sbp, ws_wata_card, ws_payments_cryptobot, ws_gifts, ws_online,
                       ws_white_counter, ws_platega_autopay, ws_platega_recurent]:
                ws.freeze_panes = ws['A2']

            fd, path = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            wb.save(path)
            return path

        export_path = await asyncio.to_thread(_sync_build_export)
        users_list = snapshot["users"]
        gifts_list = snapshot["gifts"]
        payments_list = snapshot["payments"]
        payments_cards_list = snapshot["payments_cards"]
        payments_stars_list = snapshot["payments_stars"]
        payments_platega_crypto_list = snapshot["payments_platega_crypto"]
        payments_fk_sbp_list = snapshot["payments_fk_sbp"]
        payments_wata_sbp_list = snapshot["payments_wata_sbp"]
        payments_wata_card_list = snapshot["payments_wata_card"]
        payments_cryptobot_list = snapshot["payments_cryptobot"]
        platega_autopay_list = snapshot["platega_autopay_subscriptions"]
        platega_recurent_list = snapshot["platega_recurent"]

        users_count = len(users_list)
        gifts_count = len(gifts_list)
        payments_count = len(payments_list)
        payments_cards_count = len(payments_cards_list)
        payments_stars_count = len(payments_stars_list)
        payments_cryptobot_count = len(payments_cryptobot_list)
        white_counter_count = len(snapshot["white_counter"])
        payments_platega_crypto_count = len(payments_platega_crypto_list)
        payments_fk_sbp_count = len(payments_fk_sbp_list)
        payments_wata_sbp_count = len(payments_wata_sbp_list)
        payments_wata_card_count = len(payments_wata_card_list)
        platega_autopay_count = len(platega_autopay_list)
        platega_recurent_count = len(platega_recurent_list)
        successful_payments_count = sum(1 for p in payments_list if p.status == "confirmed")
        successful_cards_count = sum(1 for p in payments_cards_list if p.status == "confirmed")
        successful_platega_crypto_count = sum(
            1 for p in payments_platega_crypto_list if p.status == "confirmed"
        )
        successful_fk_sbp_count = sum(1 for p in payments_fk_sbp_list if p.status == "confirmed")
        successful_wata_sbp_count = sum(1 for p in payments_wata_sbp_list if p.status == "confirmed")
        successful_wata_card_count = sum(1 for p in payments_wata_card_list if p.status == "confirmed")
        successful_stars_count = sum(1 for p in payments_stars_list if p.status == "confirmed")
        successful_cryptobot_count = sum(1 for p in payments_cryptobot_list if p.status == "paid")
        platega_autopay_active_count = sum(
            1 for p in platega_autopay_list if p.status in ("pending", "active", "past_due")
        )
        platega_recurent_confirmed_count = sum(
            1 for p in platega_recurent_list if p.status == "CONFIRMED"
        )

        try:
            now_s = datetime.now().strftime('%d.%m.%Y %H:%M')
            users_sheet_note = (
                "🧾 Лист <code>users</code>: все колонки таблицы.\n"
                if users_full_columns
                else ""
            )
            caption = (
                "📊 Экспорт базы данных\n"
                f"{users_sheet_note}"
                f"📅 Создано: {now_s}\n\n"
                "📊 Статистика:\n"
                f"├ 👥 Пользователей: {users_count}\n"
                f"├ 🎁 Подарков: {gifts_count}\n"
                f"├ ⚡ Платежей (legacy Platega СБП): {successful_payments_count}/{payments_count}\n"
                f"├ 💳 Платежей (legacy Platega Карта): {successful_cards_count}/{payments_cards_count}\n"
                f"├ 💳 Платежей FreeKassa (карта QR): {successful_fk_sbp_count}/{payments_fk_sbp_count}\n"
                f"├ ⚡ Платежей WATA СБП: {successful_wata_sbp_count}/{payments_wata_sbp_count}\n"
                f"├ 💳 Платежей WATA Карта: {successful_wata_card_count}/{payments_wata_card_count}\n"
                f"├ ⭐ Платежей Stars: {successful_stars_count}/{payments_stars_count}\n"
                f"├ 💰 Платежей Platega Крипто: {successful_platega_crypto_count}/{payments_platega_crypto_count}\n"
                f"├ 💎 Платежей Криптоботом: {successful_cryptobot_count}/{payments_cryptobot_count}\n"
                f"├ 🔄 Автоподписок Platega (активных): {platega_autopay_active_count}/{platega_autopay_count}\n"
                f"├ 🔄 Списаний Platega рекуррент: {platega_recurent_confirmed_count}/{platega_recurent_count}\n"
                f"└ 👁 White-кликов: {white_counter_count}"
            )
            await message.answer_document(
                document=FSInputFile(export_path),
                caption=caption,
                parse_mode="HTML",
            )
        finally:
            try:
                os.remove(export_path)
            except OSError:
                pass

        suffix = " (export_full)" if users_full_columns else ""
        logger.info(f"Администратор {message.from_user.id} экспортировал базу данных в Excel{suffix}")

    except Exception as e:
        error_message = f"❌ Ошибка при экспорте базы данных: {str(e)}"
        logger.error(error_message)
        logger.exception("Детали ошибки:")
        await message.answer(error_message)


@router.message(Command(commands=['export']))
async def export_database_to_excel(message: Message):
    """Экспорт базы данных в Excel файл."""
    await _export_database_to_excel_impl(message, users_full_columns=False)


@router.message(Command(commands=['export_full']))
async def export_full_database_to_excel(message: Message):
    """Как /export, но лист users со всеми колонками таблицы users."""
    await _export_database_to_excel_impl(message, users_full_columns=True)


MSK = ZoneInfo("Europe/Moscow")

_TRAFFIC_STAT_SNAPSHOT = date(2026, 8, 10)
_TRAFFIC_STAT_MIN_REMAINING_DAYS = 14
_TRAFFIC_STAT_TRAFFIC_GB = 7
_TRAFFIC_STAT_GREEN = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
_TRAFFIC_STAT_DATE_LABEL = "10.08"
_TRAFFIC_STAT_YES_CB = "trafic_stat_yes"
_TRAFFIC_STAT_NO_CB = "trafic_stat_no"
_TRAFFIC_STAT_PROGRESS_EVERY = 50
_TRAFFIC_STAT_PREVIEW_GB = 10
_TRAFFIC_STAT_PENDING: dict[int, list[tuple[int, int, float]]] = {}
_TRAFFIC_STAT_RUNNING: set[int] = set()
_TRAFFIC_STAT_CONFIRM_KB = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(
                text="Да, разослать всем",
                callback_data=_TRAFFIC_STAT_YES_CB,
                style=STYLE_SUCCESS,
            ),
            InlineKeyboardButton(
                text="Нет",
                callback_data=_TRAFFIC_STAT_NO_CB,
                style=STYLE_DANGER,
            ),
        ]
    ]
)


def _utc_naive(dt: datetime) -> datetime:
    if dt.tzinfo is not None:
        return dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def _payment_msk_date(utc_naive: datetime) -> date:
    return utc_naive.replace(tzinfo=timezone.utc).astimezone(MSK).date()


def _payload_map(payload: Optional[str]) -> dict[str, str]:
    if not payload:
        return {}
    out: dict[str, str] = {}
    for part in payload.split(","):
        if ":" not in part:
            continue
        k, _, v = part.partition(":")
        out[k.strip()] = v.strip()
    return out


def _trafic_stat_amount_rub(
    amount: Any,
    payload: Optional[str],
    channel: str,
    currency: Optional[str],
) -> Optional[int]:
    from handlers.handlers_statistic import convert_stars_to_rub

    try:
        raw = float(amount)
    except (TypeError, ValueError):
        raw = None

    if channel == "stars":
        if raw is None:
            return None
        mapped = convert_stars_to_rub(int(round(raw)))
        return mapped if mapped is not None else int(round(raw))

    if channel == "cryptobot" and currency and currency.upper() != "RUB":
        try:
            return int(round(float(_payload_map(payload).get("amount", ""))))
        except (TypeError, ValueError):
            if raw is None:
                return None
            return int(round(raw))

    if raw is None:
        return None
    return int(round(raw))


def _classify_trafic_stat_payment(
    payload: Optional[str],
    is_gift: bool,
    amount: Any,
    channel: str,
    currency: Optional[str],
) -> Optional[dict]:
    if is_gift:
        return None
    amount_rub = _trafic_stat_amount_rub(amount, payload, channel, currency)
    if amount_rub == 1:
        return None

    m = _payload_map(payload)
    if m.get("gift", "False").lower() == "true":
        return None

    raw_duration = m.get("duration")
    traffic_gb = _parse_traffic_duration(raw_duration)
    if traffic_gb is not None:
        return {
            "kind": "traffic",
            "days": None,
            "gb": traffic_gb,
            "amount_rub": amount_rub,
            "white": False,
        }

    white = m.get("white", "False").lower() == "true"
    days = _payload_duration_to_panel_days(raw_duration)
    if days is None and amount_rub is not None:
        days = _billing_duration_from_amount_fallback(amount_rub)
    if days is None:
        return None
    return {
        "kind": "subscription",
        "days": days,
        "gb": None,
        "amount_rub": amount_rub,
        "white": white,
    }


def _format_trafic_stat_purchase(pay_date: date, item: dict) -> str:
    ds = pay_date.strftime("%d.%m.%y")
    rub = item.get("amount_rub")
    rub_s = "—" if rub is None else f"{rub} руб"
    if item["kind"] == "traffic":
        return f"{ds} - трафик {item['gb']} ГБ - {rub_s}"
    return f"{ds} - подписка {item['days']} дней - {rub_s}"


def _naive_dt(dt: Optional[datetime]) -> Optional[datetime]:
    if dt is None:
        return None
    if dt.tzinfo is not None:
        return dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def _simulate_end_on_snapshot(
    sub_pays: List[Tuple[datetime, int]],
    snapshot: date,
) -> Optional[datetime]:
    end: Optional[datetime] = None
    ordered = sorted(
        (
            (tc, days)
            for tc, days in sub_pays
            if days and _payment_msk_date(_utc_naive(tc)) <= snapshot
        ),
        key=lambda x: _utc_naive(x[0]),
    )
    for tc, days in ordered:
        pay_t = _utc_naive(tc)
        start = end if end is not None and end > pay_t else pay_t
        end = start + timedelta(days=days)
    return end


def _end_date_as_of_snapshot(
    current_end: Optional[datetime],
    sub_pays: List[Tuple[datetime, int]],
    snapshot: date,
) -> Optional[datetime]:
    """Текущая дата окончания минус дни подписки, купленные после snapshot (МСК)."""
    current = _naive_dt(current_end)
    if current is None:
        return _simulate_end_on_snapshot(sub_pays, snapshot)

    later = sorted(
        (
            (tc, days)
            for tc, days in sub_pays
            if days and _payment_msk_date(_utc_naive(tc)) > snapshot
        ),
        key=lambda x: _utc_naive(x[0]),
        reverse=True,
    )
    end = current
    for tc, days in later:
        pay_date = _payment_msk_date(_utc_naive(tc))
        old = end - timedelta(days=days)
        if old.date() <= pay_date:
            return _simulate_end_on_snapshot(sub_pays, snapshot)
        end = old
    return end


def _build_trafic_stat_xlsx(
    users: List[Tuple[int, Optional[datetime], float, float]],
    payments: List[Tuple[int, datetime, Any, Optional[str], bool, str, Optional[str]]],
) -> Tuple[str, int, list[tuple[int, int, float]]]:
    snapshot = _TRAFFIC_STAT_SNAPSHOT
    now_msk = datetime.now(MSK).replace(tzinfo=None)

    pays_by_user: dict[int, list[tuple[datetime, dict]]] = defaultdict(list)
    sub_days_by_user: dict[int, list[tuple[datetime, int]]] = defaultdict(list)
    forever_paid: set[int] = set()
    traffic_paid: set[int] = set()

    for uid, tc, amt, pl, ig, channel, currency in payments:
        item = _classify_trafic_stat_payment(pl, ig, amt, channel, currency)
        if item is None:
            continue
        pays_by_user[uid].append((tc, item))
        if item["kind"] == "traffic":
            traffic_paid.add(uid)
            continue
        if item["kind"] == "subscription" and not item["white"] and item["days"]:
            sub_days_by_user[uid].append((tc, int(item["days"])))
            if is_forever_duration(int(item["days"])):
                forever_paid.add(uid)

    rows_out: list[tuple] = []
    apply_rows: list[tuple[int, int, float]] = []
    for uid, current_end, trafic_wl, limit_wl in users:
        if uid in forever_paid or uid in traffic_paid:
            continue
        snap_end = _end_date_as_of_snapshot(current_end, sub_days_by_user.get(uid, []), snapshot)
        if snap_end is None or snap_end.date() < snapshot:
            continue
        if is_forever_end_date(current_end) or is_forever_end_date(snap_end):
            continue
        remaining_days = (snap_end.date() - snapshot).days
        if remaining_days <= _TRAFFIC_STAT_MIN_REMAINING_DAYS:
            continue

        current = _naive_dt(current_end)
        still_active = current is not None and current > now_msk
        high_traffic = float(trafic_wl or 0) > _TRAFFIC_STAT_TRAFFIC_GB
        if not still_active or not high_traffic:
            continue
        recalc_gb = round(remaining_days / 30.0 * 10.0, 2)
        tg_id = uid
        purchases = [
            _format_trafic_stat_purchase(_payment_msk_date(_utc_naive(tc)), item)
            for tc, item in sorted(pays_by_user.get(uid, []), key=lambda x: _utc_naive(x[0]))
        ]
        apply_rows.append((uid, tg_id, recalc_gb))
        rows_out.append((
            tg_id,
            snap_end.strftime("%d.%m.%y"),
            current.strftime("%d.%m.%y") if current else "",
            "True" if still_active else None,
            trafic_wl,
            "True" if high_traffic else None,
            limit_wl,
            recalc_gb,
            purchases,
            still_active,
            high_traffic,
        ))

    rows_out.sort(key=lambda r: r[0])
    apply_rows.sort(key=lambda r: r[1] if r[1] else r[0])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "trafic_stat"

    headers = [
        "tg_id",
        f"Дата окончания на {_TRAFFIC_STAT_DATE_LABEL}",
        "Текущая дата окончания",
        "Подписка активна",
        "trafic_wl",
        "trafic_wl > 7 ГБ",
        "limit_wl",
        "Пересчитанный трафик, ГБ",
        "Покупки",
    ]
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")

    for col_num, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=title)
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_num, row in enumerate(rows_out, 2):
        (
            tg_id, snap_s, current_s, active_val, trafic_wl, traffic_val,
            limit_wl, recalc_gb, purchases, still_active, high_traffic,
        ) = row
        values = [
            tg_id,
            snap_s,
            current_s,
            active_val,
            trafic_wl,
            traffic_val,
            limit_wl,
            recalc_gb,
            "\n".join(purchases),
        ]
        n_purchases = len(purchases)
        for col_num, value in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.border = thin_border
            if col_num == 9:
                cell.alignment = wrap_top
            else:
                cell.alignment = center
            if still_active and col_num in (3, 4):
                cell.fill = _TRAFFIC_STAT_GREEN
            if high_traffic and col_num in (5, 6):
                cell.fill = _TRAFFIC_STAT_GREEN
        if n_purchases:
            ws.row_dimensions[row_num].height = min(18 * n_purchases, 180)

    widths = {
        1: 16,
        2: 24,
        3: 24,
        4: 18,
        5: 14,
        6: 18,
        7: 14,
        8: 26,
        9: 48,
    }
    for col, width in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    ws.auto_filter.ref = f"A1:I{max(1, len(rows_out) + 1)}"
    ws.freeze_panes = "A2"

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path, len(rows_out), apply_rows


def _trafic_stat_gb_label(gb: float) -> str:
    n = round(float(gb), 2)
    if n == int(n):
        return str(int(n))
    return f"{n:.2f}"


def _trafic_stat_push_text(gb: float) -> str:
    gb_s = _trafic_stat_gb_label(gb)
    return (
        "📊 ПЕРЕРАСЧЕТ ЛИМИТА ТРАФИКА\n"
        "\n"
        f"✅ Вам добавлено трафика на {gb_s} ГБ для сервера «Антиглушилка».\n"
        "\n"
        "🛡️ О сервере: Специальное решение для обхода «белых списков» операторов "
        "и борьбы с глушением VPN на мобильном интернете.\n"
        "\n"
        "📲 Что делать: Просто обновите подписку в вашем приложении, "
        "и сервер появится в общем списке для подключения.\n"
        "\n"
        "👇 Доступ уже ждет вас!"
    )


def _mirror_sub_url(raw: Optional[str]) -> str:
    if not raw:
        return ""
    return str(raw).replace(TRUE_SUB_LINK, MIRROR_SUB_LINK)


def _trafic_stat_connect_kb(sub_url: str) -> Optional[InlineKeyboardMarkup]:
    if not sub_url:
        return None
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="🔗 Подключить VPN",
                    url=sub_url,
                    style=STYLE_PRIMARY,
                )
            ]
        ]
    )


async def _trafic_stat_sub_url(
    billing_uid: int,
    panel_user: Optional[dict] = None,
) -> str:
    raw = (panel_user or {}).get("subscriptionUrl") or ""
    url = _mirror_sub_url(raw)
    if url:
        return url
    username = panel_username_for_billing_uid(billing_uid)
    return (await x3.sublink(username)) or ""


@router.message(Command(commands=["trafic_stat", "traffic_stat"]))
async def trafic_stat_excel(message: Message):
    """Активные PRO без «Навсегда»: на 10.08 > 2 недель, trafic_wl > 7, без покупок трафика."""
    if message.from_user.id not in ADMIN_IDS:
        return

    admin_id = message.from_user.id
    await message.answer("🔄 Собираю выборку /trafic_stat и формирую Excel…")
    try:
        users, payments = await sql.get_trafic_stat_source()
        path, n_rows, apply_rows = await asyncio.to_thread(
            _build_trafic_stat_xlsx, users, payments
        )
        if n_rows == 0:
            try:
                os.remove(path)
            except OSError:
                pass
            _TRAFFIC_STAT_PENDING.pop(admin_id, None)
            await message.answer(
                "Нет пользователей: активная подписка сейчас, на 10.08 дольше 2 недель, "
                "trafic_wl > 7 ГБ, без «Навсегда» и без оплат трафика."
            )
            return
        try:
            fname = f"trafic_stat_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            await message.answer_document(
                document=FSInputFile(path, filename=fname),
                caption=(
                    f"В выборке {n_rows}: активная подписка сейчас, на 10.08 дольше 2 недель, "
                    "trafic_wl > 7 ГБ, без тарифа «Навсегда», без оплат трафика. "
                    "Дата окончания на 10.08 = текущая дата в БД минус дни подписки после 10.08. "
                    "Пересчитанный трафик: (дата окончания на 10.08 − 10.08) / 30 × 10 ГБ. "
                    "Покупки: успешные платежи (подписка), не подарки."
                ),
            )
        finally:
            try:
                os.remove(path)
            except OSError:
                pass

        _TRAFFIC_STAT_PENDING[admin_id] = apply_rows
        n_push = sum(1 for _, tg_id, _gb in apply_rows if is_telegram_chat_id(tg_id))
        preview_url = await _trafic_stat_sub_url(admin_id)
        if not preview_url and apply_rows:
            preview_url = await _trafic_stat_sub_url(apply_rows[0][0])
        await message.answer("Пример пуша (10 ГБ):")
        await message.answer(
            _trafic_stat_push_text(_TRAFFIC_STAT_PREVIEW_GB),
            reply_markup=_trafic_stat_connect_kb(preview_url),
        )
        await message.answer(
            f"Разослать по выборке: <b>{n_rows}</b> чел. "
            f"(пуш уйдёт {n_push}, у кого есть Telegram id).\n"
            "Каждому: сквад с белой нодой, +пересчитанный трафик в limit_wl, пуш.\n\n"
            "Подтвердите рассылку по всем пользователям из Excel.",
            reply_markup=_TRAFFIC_STAT_CONFIRM_KB,
        )
        logger.info(f"Администратор {admin_id} выгрузил /trafic_stat ({n_rows})")
    except Exception as e:
        logger.exception("Ошибка /trafic_stat")
        await message.answer(f"❌ Ошибка при выгрузке: {e}")


@router.callback_query(F.data == _TRAFFIC_STAT_NO_CB)
async def trafic_stat_cancel(callback: CallbackQuery):
    if callback.from_user.id not in ADMIN_IDS:
        await callback.answer("Нет доступа.", show_alert=True)
        return
    _TRAFFIC_STAT_PENDING.pop(callback.from_user.id, None)
    await callback.answer()
    await callback.message.edit_text("Рассылка /trafic_stat отменена.", reply_markup=None)


@router.callback_query(F.data == _TRAFFIC_STAT_YES_CB)
async def trafic_stat_confirm(callback: CallbackQuery):
    if callback.from_user.id not in ADMIN_IDS:
        await callback.answer("Нет доступа.", show_alert=True)
        return

    admin_id = callback.from_user.id
    if admin_id in _TRAFFIC_STAT_RUNNING:
        await callback.answer("Рассылка уже идёт.", show_alert=True)
        return

    apply_rows = _TRAFFIC_STAT_PENDING.pop(admin_id, None)
    if not apply_rows:
        await callback.answer()
        await callback.message.edit_text(
            "Список пуст. Повторите /trafic_stat.",
            reply_markup=None,
        )
        return

    await callback.answer()
    total = len(apply_rows)
    await callback.message.edit_text(
        f"⏳ /trafic_stat: обработка {total} пользователей…",
        reply_markup=None,
    )

    _TRAFFIC_STAT_RUNNING.add(admin_id)
    admin_chat_id = callback.message.chat.id
    squad_moved = 0
    squad_already = 0
    squad_failed = 0
    limit_ok = 0
    limit_failed = 0
    pushed = 0
    push_failed = 0
    skipped_non_tg = 0

    try:
        for processed, (billing_uid, tg_id, recalc_gb) in enumerate(apply_rows, start=1):
            if processed % _TRAFFIC_STAT_PROGRESS_EVERY == 0:
                try:
                    await bot.send_message(
                        admin_chat_id,
                        f"trafic_stat: {processed} / {total}, "
                        f"limit {limit_ok}, squad {squad_moved}, push {pushed}",
                    )
                except Exception as notify_err:
                    logger.warning(
                        "trafic_stat: не удалось отправить прогресс админу: %s",
                        notify_err,
                    )

            try:
                await sql.add_wl_limit(billing_uid, recalc_gb)
                limit_ok += 1
            except Exception as e:
                limit_failed += 1
                logger.warning(
                    "trafic_stat: add_wl_limit uid=%s gb=%s: %s",
                    billing_uid,
                    recalc_gb,
                    e,
                )

            panel_user = await fetch_panel_user(x3, billing_uid)
            if not panel_user:
                squad_failed += 1
            elif user_on_active_squad(panel_user):
                squad_already += 1
            elif await reassign_to_active_squad(x3, panel_user):
                squad_moved += 1
            else:
                squad_failed += 1

            if not is_telegram_chat_id(tg_id):
                skipped_non_tg += 1
                await asyncio.sleep(0.05)
                continue

            try:
                sub_url = await _trafic_stat_sub_url(billing_uid, panel_user)
                await bot.send_message(
                    chat_id=tg_id,
                    text=_trafic_stat_push_text(recalc_gb),
                    reply_markup=_trafic_stat_connect_kb(sub_url),
                )
                pushed += 1
            except Exception as e:
                push_failed += 1
                logger.warning("trafic_stat: push uid=%s tg=%s: %s", billing_uid, tg_id, e)

            await asyncio.sleep(0.05)
    finally:
        _TRAFFIC_STAT_RUNNING.discard(admin_id)

    await bot.send_message(
        admin_chat_id,
        (
            "✅ <b>Готово (/trafic_stat)</b>\n\n"
            f"• В выборке: <b>{total}</b>\n"
            f"• +limit_wl: <b>{limit_ok}</b>\n"
            f"• Squad → белая нода: <b>{squad_moved}</b>\n"
            f"• Уже на белой ноде: {squad_already}\n"
            f"• Push отправлено: <b>{pushed}</b>\n"
            f"• Ошибка limit_wl: {limit_failed}\n"
            f"• Ошибка squad: {squad_failed}\n"
            f"• Ошибка push: {push_failed}\n"
            f"• Пропущено (не Telegram chat_id): {skipped_non_tg}"
        ),
    )
    logger.info(
        "Админ %s /trafic_stat apply: total=%s limit=%s squad=%s pushed=%s",
        admin_id,
        total,
        limit_ok,
        squad_moved,
        pushed,
    )


@router.message(Command("export_panel"))
async def export_panel(message: Message):
    if message.from_user.id not in ADMIN_IDS:
        return

    users_x3 = await x3.get_all_panel()
    total = len(users_x3)
    await message.answer(f"{total} - всего юзеров в панели. Формирую Excel...")

    if not users_x3:
        await message.answer("Нет пользователей для экспорта.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "panel_users"

    # Заголовки
    headers = [
        "username", "telegramId", "expireAt",
        "shortUuid", "vlessUuid", "trojanPassword", "ssPassword",
        "description", "squad_uuid"
    ]
    ws.append(headers)

    # Стили
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    # Заголовки форматируем
    for col_num, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=title)
        cell.alignment = header_alignment
        cell.border = thin_border

    # Заполнение данными
    for user in users_x3:
        # Извлекаем squad (первый элемент списка activeInternalSquads, если есть)
        squad_name = ""
        squad_uuid = ""
        if user.get('activeInternalSquads') and len(user['activeInternalSquads']) > 0:
            squad = user['activeInternalSquads'][0]
            squad_name = squad.get('name', '')
            squad_uuid = squad.get('uuid', '')

        # Форматируем даты (если есть)
        def format_date(dt_str):
            if dt_str:
                try:
                    dt = datetime.fromisoformat(dt_str.replace('Z', '+00:00'))
                    return dt.strftime('%Y-%m-%d %H:%M:%S')
                except:
                    return dt_str
            return ""

        row_data = [
            user.get('username', ''),
            user.get('telegramId', ''),
            format_date(user.get('expireAt')),
            user.get('shortUuid', ''),
            user.get('vlessUuid', ''),
            user.get('trojanPassword', ''),
            user.get('ssPassword', ''),
            user.get('description', ''),
            squad_uuid
        ]
        ws.append(row_data)

    # Автоширина колонок
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    # Заморозка заголовка
    ws.freeze_panes = 'A2'

    wb.save('panel.xlsx')

    # Отправляем файл
    from aiogram.types import BufferedInputFile
    await message.answer_document(
        document=FSInputFile('panel.xlsx',
        filename=f"panel_users_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"),
        caption=f"📊 Выгружено пользователей из панели: {total}"
    )

    logger.info(f"Администратор {message.from_user.id} выгрузил список пользователей панели")
